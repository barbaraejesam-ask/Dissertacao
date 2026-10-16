#!/usr/bin/env python3
"""
Gerador do Mapa de Indicadores de Saneamento de Santa Catarina.

Le os arquivos IN*_Resultado_Analise_SC.xlsx e a malha municipal em GeoJSON,
consolida a classificacao por municipio/ano/indicador e escreve um unico
arquivo HTML autocontido em public/index.html, pronto para deploy estatico.

Uso:
    python build.py
    python build.py --ano-min 2003 --ano-max 2024 --precisao 3
    python build.py --dados ./data --saida ./public
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Dependencia ausente. Rode: pip install openpyxl")

RAIZ = Path(__file__).resolve().parent

# Ordem de exibicao e rotulos. Chave = prefixo do nome do arquivo.
INDICADORES = {
    "IN023": "IN023 - Atendimento urbano de agua",
    "IN024": "IN024 - Atendimento urbano de esgoto",
    "IN016": "IN016 - Tratamento de esgoto",
    "IN084": "IN084 - Coliformes fora do padrao",
    "IN082": "IN082 - Extravasamentos de esgoto",
    "IN049": "IN049 - Perdas na distribuicao",
    "IN009": "IN009 - Indice de hidrometracao",
    "IN011": "IN011 - Indice de macromedicao",
}

# Indicadores em que valor MENOR representa desempenho melhor.
# Usado apenas para rotular o sentido no HTML; a classificacao vem pronta do xlsx.
MENOR_MELHOR = {"IN049", "IN082", "IN084"}

CODIGO_CLASSE = {
    "Satisfatório": 1,
    "Satisfatorio": 1,
    "Moderado": 2,
    "Insatisfatório": 3,
    "Insatisfatorio": 3,
}


def norm(texto) -> str:
    """Normaliza nome de coluna para comparacao tolerante."""
    return re.sub(r"[^a-z0-9]", "", str(texto).lower()) if texto is not None else ""


def achar_coluna(cabecalho: list, *candidatos: str) -> int:
    """Encontra o indice da primeira coluna que casa com um dos candidatos."""
    normalizado = [norm(c) for c in cabecalho]
    for cand in candidatos:
        alvo = norm(cand)
        if alvo in normalizado:
            return normalizado.index(alvo)
    raise KeyError(
        f"Coluna nao encontrada. Procurei por {candidatos}, cabecalho tem {cabecalho}"
    )


def ler_indicador(caminho: Path, sigla: str, ano_min: int, ano_max: int) -> dict:
    """Le a primeira aba do xlsx e retorna {ano: {cod_ibge: [classe, valor]}}."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = list(next(linhas))

    i_cod = achar_coluna(cabecalho, "CD_MUN", "cod_ibge", "codigo_municipio")
    i_ano = achar_coluna(cabecalho, "ano", "ano_referencia")
    i_val = achar_coluna(cabecalho, sigla)
    i_cls = achar_coluna(cabecalho, "classificacao", "classificação", "status")

    saida: dict[int, dict[str, list]] = {}
    descartadas = 0

    for linha in linhas:
        cod, ano = linha[i_cod], linha[i_ano]
        if cod is None or ano is None:
            descartadas += 1
            continue
        ano = int(ano)
        if not (ano_min <= ano <= ano_max):
            continue
        valor = linha[i_val]
        valor = round(float(valor), 2) if isinstance(valor, (int, float)) else None
        classe = CODIGO_CLASSE.get(str(linha[i_cls]).strip(), 0)
        saida.setdefault(ano, {})[str(int(cod))] = [classe, valor]

    wb.close()
    if descartadas:
        print(f"  aviso: {descartadas} linha(s) sem codigo ou ano em {caminho.name}")
    return saida


def simplificar_geojson(caminho: Path, precisao: int) -> dict:
    """Carrega o GeoJSON, arredonda coordenadas e mantem so id e name."""
    geo = json.loads(caminho.read_text(encoding="utf-8"))

    def arredonda(coords):
        if coords and isinstance(coords[0], (int, float)):
            return [round(coords[0], precisao), round(coords[1], precisao)]
        return [arredonda(c) for c in coords]

    for feicao in geo["features"]:
        feicao["geometry"]["coordinates"] = arredonda(feicao["geometry"]["coordinates"])
        props = feicao["properties"]
        ident = str(props.get("id") or props.get("CD_MUN") or props.get("codarea") or "")
        nome = props.get("name") or props.get("NM_MUN") or props.get("nome") or ident
        feicao["properties"] = {"id": ident, "name": nome}

    return geo


def conferir_cobertura(geo: dict, dados: dict) -> None:
    """Avisa sobre municipios presentes nos dados mas ausentes da malha."""
    na_malha = {f["properties"]["id"] for f in geo["features"]}
    nos_dados = {
        cod
        for anos in dados.values()
        for mapa in anos.values()
        for cod in mapa
    }
    faltando = sorted(nos_dados - na_malha)
    if faltando:
        print(
            f"  aviso: {len(faltando)} municipio(s) nos dados sem poligono na malha "
            f"(nao serao pintados): {', '.join(faltando[:8])}"
            + (" ..." if len(faltando) > 8 else "")
        )
    print(f"  malha com {len(na_malha)} poligonos, dados com {len(nos_dados)} municipios")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dados", default=str(RAIZ / "data"), help="pasta com xlsx e geojson")
    ap.add_argument("--saida", default=str(RAIZ / "public"), help="pasta de saida")
    ap.add_argument("--template", default=str(RAIZ / "src" / "template.html"))
    ap.add_argument("--geojson", default="sc_municipios.geojson")
    ap.add_argument("--ano-min", type=int, default=2003)
    ap.add_argument("--ano-max", type=int, default=2024)
    ap.add_argument("--precisao", type=int, default=3, help="casas decimais das coordenadas")
    ap.add_argument("--csv", action="store_true", help="exportar tambem o CSV consolidado")
    args = ap.parse_args()

    pasta_dados = Path(args.dados)
    pasta_saida = Path(args.saida)
    pasta_saida.mkdir(parents=True, exist_ok=True)

    print("1. Lendo planilhas de indicadores")
    dados: dict[str, dict] = {}
    for sigla in INDICADORES:
        arquivos = list(pasta_dados.glob(f"{sigla}*.xlsx"))
        if not arquivos:
            print(f"  {sigla}: arquivo nao encontrado, pulando")
            continue
        dados[sigla] = ler_indicador(arquivos[0], sigla, args.ano_min, args.ano_max)
        n_regs = sum(len(v) for v in dados[sigla].values())
        print(f"  {sigla}: {len(dados[sigla])} anos, {n_regs} registros")

    if not dados:
        sys.exit("Nenhum indicador carregado. Confira a pasta de dados.")

    print("2. Carregando e simplificando a malha municipal")
    geo = simplificar_geojson(pasta_dados / args.geojson, args.precisao)
    conferir_cobertura(geo, dados)

    print("3. Montando o HTML")
    template = Path(args.template).read_text(encoding="utf-8")
    pacote = {
        "inds": list(dados.keys()),
        "labels": INDICADORES,
        "menorMelhor": sorted(MENOR_MELHOR),
        "anoMin": args.ano_min,
        "anoMax": args.ano_max,
        "data": dados,
    }
    html = template.replace(
        "__GEOJSON__", json.dumps(geo, ensure_ascii=False, separators=(",", ":"))
    ).replace(
        "__DADOS__", json.dumps(pacote, ensure_ascii=False, separators=(",", ":"))
    )

    destino = pasta_saida / "index.html"
    destino.write_text(html, encoding="utf-8")
    print(f"  gerado {destino} ({destino.stat().st_size / 1_048_576:.2f} MB)")

    if args.csv:
        print("4. Exportando CSV consolidado (formato longo)")
        nomes = {f["properties"]["id"]: f["properties"]["name"] for f in geo["features"]}
        linhas = ["cod_ibge,municipio,ano,indicador,valor,classificacao"]
        rotulo = {0: "Sem dados", 1: "Satisfatorio", 2: "Moderado", 3: "Insatisfatorio"}
        for sigla, anos in dados.items():
            for ano in sorted(anos):
                for cod, (classe, valor) in sorted(anos[ano].items()):
                    nome = str(nomes.get(cod, "")).replace(",", " ")
                    linhas.append(
                        f"{cod},{nome},{ano},{sigla},"
                        f"{'' if valor is None else valor},{rotulo[classe]}"
                    )
        csv_path = pasta_saida / "indicadores_sc_long.csv"
        csv_path.write_text("\n".join(linhas), encoding="utf-8")
        print(f"  gerado {csv_path} ({len(linhas) - 1} linhas)")

    print("Concluido.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
